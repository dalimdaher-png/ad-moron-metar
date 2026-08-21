#!/usr/bin/env python3
"""Exporta el cruce METAR/SPECI -> IMC/VMC de un mes a JSON (con movimientos
y marca manual incluidos), para alimentar el visor interactivo hora por hora.

Uso:
    python exportar_datos.py "datos/PLANILLA MOV ACFT AGOSTO 2026 - original.xlsx" datos/sadm_agosto.csv datos_horarios.json
"""
import csv, json, re, sys
from datetime import date, datetime, timedelta
from openpyxl import load_workbook
from metar_moron import clasificar, VIS_ILIM, TECHO_ILIM

VIS_MIN, TECHO_MIN = 5000, 1500
FILA0, FILAN = 10, 33

TOKEN_OK = re.compile(
    r'^(?:'
    r'(?:VRB|\d{3})\d{2,3}(?:G\d{2,3})?(?:KT|MPS|KMH)|\d{3}V\d{3}|'
    r'\d{4}[NESW]{0,2}|\d{1,2}SM|M?\d/\dSM|'
    r'(?:FEW|SCT|BKN|OVC)(?:\d{3}|///)(?:CB|TCU)?|VV(?:\d{3}|///)|'
    r'NSC|NCD|SKC|CLR|CAVOK|COR|AMD|RTD|AUTO|NIL|RE\w*|WS\w*|'
    r'M?\d{2}/M?\d{2}|/{2}//{0,2}|Q\d{4}|A\d{4}|'
    r'[-+]?(?:VC)?(?:MI|BC|PR|DR|BL|SH|TS|FZ|RE)?'
    r'(?:DZ|RA|SN|SG|PL|GR|GS|UP|BR|FG|FU|VA|DU|SA|HZ|PO|SQ|FC|SS|DS)+'
    r')$')

def raros(raw):
    t = raw.strip().rstrip('=').split()
    ini = next((i + 1 for i, x in enumerate(t) if re.fullmatch(r'\d{6}Z', x)), 0)
    cuerpo = t[ini:]
    for stop in ('TEMPO', 'BECMG', 'NOSIG', 'RMK'):
        if stop in cuerpo: cuerpo = cuerpo[:cuerpo.index(stop)]
    return [x for x in cuerpo if not TOKEN_OK.match(x)]

def cargar(ruta):
    out = []
    for r in csv.reader(open(ruta)):
        if len(r) >= 3 and r[0] != 'station' and r[2] not in ('M', ''):
            out.append((datetime.strptime(r[1], '%Y-%m-%d %H:%M'), r[2]))
    return sorted(out)

def por_hora(reportes, d1, d2):
    idx = {}
    for dt, raw in reportes:
        idx.setdefault(dt.replace(minute=0), []).append((dt, raw))
    filas, cur = [], datetime.combine(d1, datetime.min.time())
    fin = datetime.combine(d2, datetime.min.time()) + timedelta(days=1)
    while cur < fin:
        g = idx.get(cur, [])
        cl = [(dt, raw, clasificar(raw, VIS_MIN, TECHO_MIN)) for dt, raw in g]
        ok = [c for c in cl if c[2]['cond'] != 'SIN DATO']
        if not ok:
            filas.append(dict(dt=cur, cond='SIN DATO', vis=None, techo=None, speci='NO',
                              metar='', todos='', motivo='', raros=[]))
        else:
            base = next((c for c in ok if c[0].minute == 0), ok[0])
            peor = min(ok, key=lambda c: (c[2]['vis_m'], c[2]['techo_ft']))
            filas.append(dict(
                dt=cur,
                cond='IMC' if any(c[2]['cond'] == 'IMC' for c in ok) else 'VMC',
                vis=min(c[2]['vis_m'] for c in ok),
                techo=min(c[2]['techo_ft'] for c in ok),
                speci='SI' if any(c[0].minute != 0 for c in cl) else 'NO',
                metar=base[1],
                todos=' || '.join(r for _, r, _ in cl),
                motivo=peor[2]['motivo'],
                raros=sorted({t for _, r, _ in cl for t in raros(r)})))
        cur += timedelta(hours=1)
    return filas

def main():
    original, csvf, salida = sys.argv[1], sys.argv[2], sys.argv[3]
    wbv = load_workbook(original, data_only=True)
    reps = cargar(csvf)
    d1, d2 = reps[0][0].date(), reps[-1][0].date()
    horas = por_hora(reps, d1, d2)
    mapa = {(h['dt'].date(), h['dt'].hour): h for h in horas}

    dias = {}
    for n in wbv.sheetnames:
        if not n.isdigit(): continue
        ws = wbv[n]
        f = ws['I6'].value
        fecha = f.date() if isinstance(f, datetime) else f
        if not isinstance(fecha, date): continue
        registros = []
        num = lambda v: v if isinstance(v, (int, float)) else 0
        for i, fila in enumerate(range(FILA0, FILAN + 1)):
            manual = any(isinstance(ws[f'{c}{fila}'].value, str) and 'IMC' in ws[f'{c}{fila}'].value.upper()
                         for c in 'CDEFGHIJKLM')
            mv = ws[f'H{fila}'].value
            mv = mv if isinstance(mv, (int, float)) else 0
            mov_detalle = {
                'comercial': num(ws[f'C{fila}'].value),
                'oficial': num(ws[f'D{fila}'].value),
                'otros': num(ws[f'E{fila}'].value),
                'local': num(ws[f'F{fila}'].value),
                'sobrevuelo': num(ws[f'G{fila}'].value),
                'ifr': num(ws[f'I{fila}'].value),
            }
            h = mapa.get((fecha, i))
            registros.append({
                'hora': i,
                'condicion': h['cond'] if h else 'SIN DATO',
                'vis': None if not h or h['vis'] is None or h['vis'] == VIS_ILIM else h['vis'],
                'vis_ilim': bool(h and h['vis'] == VIS_ILIM),
                'techo': None if not h or h['techo'] is None or h['techo'] == TECHO_ILIM else h['techo'],
                'techo_ilim': bool(h and h['techo'] == TECHO_ILIM),
                'speci': bool(h and h['speci'] == 'SI'),
                'metar': (h['metar'] if h else '') or '',
                'todos': (h['todos'] if h else '') or '',
                'motivo': (h['motivo'] if h else '') or '',
                'raros': (h['raros'] if h else []) or [],
                'movimientos': mv,
                'mov_detalle': mov_detalle,
                'marca_manual': manual,
            })
        dias[str(fecha)] = registros

    with open(salida, 'w', encoding='utf-8') as f:
        json.dump({'estacion': 'SADM', 'desde': str(d1), 'hasta': str(d2),
                    'generado': datetime.now().isoformat(timespec='minutes'), 'dias': dias},
                   f, ensure_ascii=False)
    print(f"{salida}: {len(dias)} dias exportados")

if __name__ == '__main__':
    main()
