#!/usr/bin/env python3
"""
Cruce METAR/SPECI -> IMC/VMC para SADM (Moron, Argentina).

Uso:
    python metar_moron.py 2026-08-07 2026-08-21
    python metar_moron.py 2026-08-07 2026-08-21 --estacion SADM --vis 5000 --techo 1500

Genera meteo_<estacion>_<desde>_<hasta>.xlsx con 4 hojas listas para BUSCARV.
Fuente: Iowa Environmental Mesonet (archivo publico de METAR/SPECI, mismo texto crudo que Ogimet).
"""
import argparse, csv, io, re, sys
from datetime import datetime, timedelta, date

# ------------------------- umbrales -------------------------
VIS_MIN_M    = 5000    # visibilidad minima para VMC (metros)
TECHO_MIN_FT = 1500    # techo minimo para VMC (pies AGL)
VIS_ILIM     = 10000   # 9999 / CAVOK
TECHO_ILIM   = 99999   # sin capa BKN/OVC/VV

FENOM = (r'DZ|RA|SN|SG|PL|GR|GS|UP|BR|FG|FU|VA|DU|SA|HZ|PO|SQ|FC|SS|DS')
RE_FENOM = re.compile(r'^[-+]?(VC)?(MI|BC|PR|DR|BL|SH|TS|FZ|RE)?(' + FENOM + r')+$')

# ------------------------- parser -------------------------
def parse_metar(raw):
    t = raw.strip().rstrip('=')
    toks = t.split()
    out = {'vis_m': None, 'techo_ft': None, 'cavok': False, 'fenom': [], 'nil': False}

    # arrancar despues del grupo de hora (ddhhmmZ); saltea COR/AMD/RTD y el indicativo
    ini = 0
    for i, tk in enumerate(toks):
        if re.fullmatch(r'\d{6}Z', tk):
            ini = i + 1
            break
    toks = toks[ini:]

    if 'NIL' in toks:
        out['nil'] = True
        return out
    # cortar en tendencia / remarks: eso es pronostico, no estado actual
    for stop in ('TEMPO', 'BECMG', 'NOSIG', 'RMK'):
        if stop in toks:
            toks = toks[:toks.index(stop)]
    if 'CAVOK' in toks:
        return {**out, 'cavok': True, 'vis_m': VIS_ILIM, 'techo_ft': TECHO_ILIM}

    capas, nsc = [], False
    for tk in toks:
        if re.fullmatch(r'(VRB|\d{3})\d{2,3}(G\d{2,3})?(KT|MPS|KMH)', tk):  # viento
            continue
        if re.fullmatch(r'\d{3}V\d{3}', tk):                                # variacion viento
            continue
        m = re.fullmatch(r'(\d{4})(N|NE|E|SE|S|SW|W|NW)?', tk)              # vis en metros
        if m and out['vis_m'] is None:
            v = int(m.group(1)); out['vis_m'] = VIS_ILIM if v == 9999 else v; continue
        m = re.fullmatch(r'(\d+)(?:\s)?(?:(\d+)/(\d+))?SM', tk)             # vis en millas
        if m and out['vis_m'] is None:
            out['vis_m'] = int(int(m.group(1)) * 1609.344); continue
        m = re.fullmatch(r'(FEW|SCT|BKN|OVC)(\d{3}|///)(CB|TCU)?', tk)      # capas
        if m:
            if m.group(2) != '///':
                capas.append((m.group(1), int(m.group(2)) * 100))
            continue
        m = re.fullmatch(r'VV(\d{3}|///)', tk)                              # vis vertical
        if m:
            capas.append(('VV', int(m.group(1)) * 100 if m.group(1) != '///' else 0)); continue
        if tk in ('NSC', 'NCD', 'SKC', 'CLR'):
            nsc = True; continue
        if RE_FENOM.match(tk):
            out['fenom'].append(tk); continue

    techos = [alt for tipo, alt in capas if tipo in ('BKN', 'OVC', 'VV')]
    out['techo_ft'] = min(techos) if techos else TECHO_ILIM
    if out['vis_m'] is None and (nsc or capas):
        out['vis_m'] = VIS_ILIM
    return out

def clasificar(raw, vis_min=VIS_MIN_M, techo_min=TECHO_MIN_FT):
    p = parse_metar(raw)
    if p['nil'] or p['vis_m'] is None:
        return {**p, 'cond': 'SIN DATO', 'motivo': ''}
    mot = []
    if p['vis_m'] < vis_min:    mot.append(f"vis {p['vis_m']} m")
    if p['techo_ft'] < techo_min: mot.append(f"techo {p['techo_ft']} ft")
    return {**p, 'cond': 'IMC' if mot else 'VMC', 'motivo': ' + '.join(mot)}

# ------------------------- descarga -------------------------
IEM = ("https://mesonet.agron.iastate.edu/cgi-bin/request/asos.py"
       "?station={est}&data=metar&year1={a1}&month1={m1}&day1={d1}"
       "&year2={a2}&month2={m2}&day2={d2}&tz=Etc/UTC&format=onlycomma"
       "&missing=M&trace=T&direct=no")

def descargar(est, desde, hasta, pausa=3.0):
    """Baja en tramos de 7 dias para no golpear el servicio (devuelve 429 si se abusa)."""
    import time, urllib.request
    filas, cur = [], desde
    while cur <= hasta:
        fin = min(cur + timedelta(days=7), hasta + timedelta(days=1))
        url = IEM.format(est=est, a1=cur.year, m1=cur.month, d1=cur.day,
                         a2=fin.year, m2=fin.month, d2=fin.day)
        for intento in range(4):
            try:
                with urllib.request.urlopen(url, timeout=120) as r:
                    txt = r.read().decode('utf-8', 'replace')
                break
            except Exception as e:
                if intento == 3: raise
                print(f"  reintento {intento+1} ({e})", file=sys.stderr); time.sleep(15)
        for row in csv.reader(io.StringIO(txt)):
            if len(row) >= 3 and row[0] != 'station' and row[2] not in ('M', ''):
                filas.append((row[1], row[2]))
        print(f"  {cur} -> {fin}: {len(filas)} reportes acumulados", file=sys.stderr)
        cur = fin; time.sleep(pausa)
    return filas

# ------------------------- armado -------------------------
def construir(filas, desde, hasta, est, vis_min, techo_min):
    rep = []
    for ts, raw in filas:
        dt = datetime.strptime(ts, '%Y-%m-%d %H:%M')
        r = clasificar(raw, vis_min, techo_min)
        rep.append({'dt': dt, 'tipo': 'METAR' if dt.minute == 0 else 'SPECI', 'raw': raw, **r})
    rep.sort(key=lambda x: x['dt'])

    por_hora, cur = [], datetime.combine(desde, datetime.min.time())
    fin = datetime.combine(hasta, datetime.min.time()) + timedelta(days=1)
    idx = {}
    for r in rep:
        idx.setdefault(r['dt'].replace(minute=0), []).append(r)
    while cur < fin:
        grupo = idx.get(cur, [])
        validos = [g for g in grupo if g['cond'] != 'SIN DATO']
        if not validos:
            por_hora.append({'dt': cur, 'cond': 'SIN DATO', 'vis': None, 'techo': None,
                             'fenom': '', 'speci': 'NO', 'metar': '', 'detalle': '', 'motivo': ''})
        else:
            peor = min(validos, key=lambda g: (g['vis_m'], g['techo_ft']))
            base = next((g for g in validos if g['tipo'] == 'METAR'), validos[0])
            por_hora.append({
                'dt': cur,
                'cond': 'IMC' if any(g['cond'] == 'IMC' for g in validos) else 'VMC',
                'vis': min(g['vis_m'] for g in validos),
                'techo': min(g['techo_ft'] for g in validos),
                'fenom': ' '.join(dict.fromkeys(sum((g['fenom'] for g in validos), []))),
                'speci': 'SI' if any(g['tipo'] == 'SPECI' for g in validos) else 'NO',
                'metar': base['raw'],
                'detalle': ' || '.join(g['raw'] for g in grupo),
                'motivo': peor['motivo'],
            })
        cur += timedelta(hours=1)
    return rep, por_hora

def escribir(rep, por_hora, ruta, est, vis_min, techo_min):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    ILIM = lambda v, lim: ('ILIM' if v == lim else v) if v is not None else 'S/D'
    wb = Workbook()
    hdr_f  = Font(bold=True, color='FFFFFF')
    hdr_fl = PatternFill('solid', fgColor='1F3864')
    imc    = PatternFill('solid', fgColor='F8CBAD')
    sd     = PatternFill('solid', fgColor='D9D9D9')

    def hoja(ws, cols, filas, anchos, col_cond=None):
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = hdr_f, hdr_fl
            c.alignment = Alignment(horizontal='center')
        for f in filas: ws.append(f)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for i, w in enumerate(anchos, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        if col_cond:
            for r in range(2, ws.max_row + 1):
                v = ws.cell(r, col_cond).value
                if v == 'IMC':
                    for c in range(1, len(cols) + 1): ws.cell(r, c).fill = imc
                elif v == 'SIN DATO':
                    for c in range(1, len(cols) + 1): ws.cell(r, c).fill = sd

    ws = wb.active; ws.title = 'POR_HORA'
    hoja(ws,
         ['CLAVE', 'FECHA', 'HORA_Z', 'CONDICION', 'MOTIVO', 'VIS_M', 'TECHO_FT',
          'FENOMENO', 'HUBO_SPECI', 'METAR', 'TODOS_LOS_REPORTES'],
         [[h['dt'].strftime('%Y-%m-%d %H'), h['dt'].strftime('%Y-%m-%d'), h['dt'].hour,
           h['cond'], h['motivo'], ILIM(h['vis'], VIS_ILIM), ILIM(h['techo'], TECHO_ILIM),
           h['fenom'], h['speci'], h['metar'], h['detalle']] for h in por_hora],
         [15, 12, 8, 11, 22, 9, 10, 12, 11, 60, 80], col_cond=4)

    ws = wb.create_sheet('REPORTES')
    hoja(ws,
         ['CLAVE', 'FECHA_HORA_Z', 'TIPO', 'CONDICION', 'MOTIVO', 'VIS_M', 'TECHO_FT', 'FENOMENO', 'METAR'],
         [[r['dt'].strftime('%Y-%m-%d %H'), r['dt'].strftime('%Y-%m-%d %H:%M'), r['tipo'],
           r['cond'], r['motivo'], ILIM(r['vis_m'], VIS_ILIM), ILIM(r['techo_ft'], TECHO_ILIM),
           ' '.join(r['fenom']), r['raw']] for r in rep],
         [15, 18, 8, 11, 22, 9, 10, 12, 60], col_cond=4)

    ws = wb.create_sheet('RESUMEN_DIARIO')
    dias = {}
    for h in por_hora:
        d = dias.setdefault(h['dt'].date(), {'imc': 0, 'vmc': 0, 'sd': 0, 'vis': [], 'techo': [], 'fen': set()})
        if h['cond'] == 'IMC':   d['imc'] += 1
        elif h['cond'] == 'VMC': d['vmc'] += 1
        else:                    d['sd'] += 1
        if h['vis'] is not None:   d['vis'].append(h['vis'])
        if h['techo'] is not None: d['techo'].append(h['techo'])
        if h['fenom']: d['fen'].update(h['fenom'].split())
    hoja(ws,
         ['FECHA', 'HS_IMC', 'HS_VMC', 'HS_SIN_DATO', '%_IMC', 'PEOR_VIS_M', 'PEOR_TECHO_FT', 'FENOMENOS'],
         [[str(d), v['imc'], v['vmc'], v['sd'],
           round(100 * v['imc'] / max(1, v['imc'] + v['vmc']), 1),
           ILIM(min(v['vis']) if v['vis'] else None, VIS_ILIM),
           ILIM(min(v['techo']) if v['techo'] else None, TECHO_ILIM),
           ' '.join(sorted(v['fen']))] for d, v in sorted(dias.items())],
         [12, 9, 9, 13, 9, 12, 14, 20])

    ws = wb.create_sheet('PARAMETROS')
    for f in [['Estacion', est],
              ['Visibilidad minima VMC (m)', vis_min],
              ['Techo minimo VMC (ft)', techo_min],
              ['Regla', 'IMC si visibilidad < umbral O techo < umbral'],
              ['Techo', 'capa mas baja BKN/OVC/VV. SCT y FEW NO son techo.'],
              ['CAVOK', 'vis >= 10 km y sin nubes significativas -> VMC'],
              ['SPECI', 'incluidos. CONDICION de la hora = la peor condicion reportada dentro de esa hora.'],
              ['Huso', 'TODO EN UTC (Z), igual que la planilla de movimientos'],
              ['Fuente', 'Iowa Environmental Mesonet - archivo publico de METAR/SPECI'],
              ['Generado', datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')]]:
        ws.append(f)
    ws.column_dimensions['A'].width = 30; ws.column_dimensions['B'].width = 85
    for r in range(1, ws.max_row + 1): ws.cell(r, 1).font = Font(bold=True)

    wb.save(ruta)
    return ruta

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('desde'); ap.add_argument('hasta')
    ap.add_argument('--estacion', default='SADM')
    ap.add_argument('--vis', type=int, default=VIS_MIN_M)
    ap.add_argument('--techo', type=int, default=TECHO_MIN_FT)
    ap.add_argument('--csv', help='usar un CSV ya descargado en vez de bajar de internet')
    a = ap.parse_args()
    d1 = date.fromisoformat(a.desde); d2 = date.fromisoformat(a.hasta)
    if a.csv:
        filas = [(r[1], r[2]) for r in csv.reader(open(a.csv))
                 if len(r) >= 3 and r[0] != 'station' and r[2] not in ('M', '')]
    else:
        filas = descargar(a.estacion, d1, d2)
    rep, ph = construir(filas, d1, d2, a.estacion, a.vis, a.techo)
    ruta = f"meteo_{a.estacion}_{a.desde}_{a.hasta}.xlsx"
    escribir(rep, ph, ruta, a.estacion, a.vis, a.techo)
    n_imc = sum(1 for h in ph if h['cond'] == 'IMC')
    print(f"{ruta}: {len(rep)} reportes, {len(ph)} horas, {n_imc} horas IMC")

if __name__ == '__main__':
    main()
