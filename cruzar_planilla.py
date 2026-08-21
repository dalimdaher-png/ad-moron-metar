#!/usr/bin/env python3
"""Inyecta el cruce METAR/SPECI -> IMC/VMC dentro de la PLANILLA MOV ACFT, hora por hora."""
import csv, re, sys, shutil
from datetime import datetime, timedelta, date
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from metar_moron import clasificar, parse_metar, VIS_ILIM, TECHO_ILIM

VIS_MIN, TECHO_MIN = 5000, 1500
FILA0, FILAN = 10, 33          # franjas 00a01 .. 23a00
COLS = ['R', 'S', 'T', 'U', 'V']
CAB  = ['CONDICION', 'VIS (m)', 'TECHO (ft)', 'SPECI', 'METAR OBSERVADO']

TOKEN_OK = re.compile(
    r'^(?:'
    r'(?:VRB|\d{3})\d{2,3}(?:G\d{2,3})?(?:KT|MPS|KMH)|\d{3}V\d{3}|'
    r'\d{4}[NESW]{0,2}|\d{1,2}SM|M?\d/\dSM|'
    r'(?:FEW|SCT|BKN|OVC)(?:\d{3}|///)(?:CB|TCU)?|VV(?:\d{3}|///)|'
    r'NSC|NCD|SKC|CLR|CAVOK|COR|AMD|RTD|AUTO|NIL|RE\w*|WS\w*|'
    r'M?\d{2}/M?\d{2}|/{2}//{0,2}|Q\d{4}|A\d{4}|'
    r'[-+]?(?:VC)?(?:MI|BC|PR|DR|BL|SH|TS|FZ|RE)?'
    r'(?:DZ|RA|SN|SG|PL|GR|GS|UP|BR|FG|FU|VA|DU|SA|HZ|PO|SQ|FC|SS|DS)+'
    r')$')

def raros(raw):
    t = raw.strip().rstrip('=').split()
    ini = next((i + 1 for i, x in enumerate(t) if re.fullmatch(r'\d{6}Z', x)), 0)
    cuerpo = t[ini:]
    for stop in ('TEMPO', 'BECMG', 'NOSIG', 'RMK'):
        if stop in cuerpo: cuerpo = cuerpo[:cuerpo.index(stop)]
    return [x for x in cuerpo if not TOKEN_OK.match(x)]

def cargar(ruta):
    out = []
    for r in csv.reader(open(ruta)):
        if len(r) >= 3 and r[0] != 'station' and r[2] not in ('M', ''):
            out.append((datetime.strptime(r[1], '%Y-%m-%d %H:%M'), r[2]))
    return sorted(out)

def por_hora(reportes, d1, d2):
    idx = {}
    for dt, raw in reportes:
        idx.setdefault(dt.replace(minute=0), []).append((dt, raw))
    filas, cur = [], datetime.combine(d1, datetime.min.time())
    fin = datetime.combine(d2, datetime.min.time()) + timedelta(days=1)
    while cur < fin:
        g = idx.get(cur, [])
        cl = [(dt, raw, clasificar(raw, VIS_MIN, TECHO_MIN)) for dt, raw in g]
        ok = [c for c in cl if c[2]['cond'] != 'SIN DATO']
        if not ok:
            filas.append(dict(dt=cur, cond='SIN DATO', vis=None, techo=None, speci='NO',
                              metar='', todos='', motivo='', raros=[]))
        else:
            base = next((c for c in ok if c[0].minute == 0), ok[0])
            peor = min(ok, key=lambda c: (c[2]['vis_m'], c[2]['techo_ft']))
            filas.append(dict(
                dt=cur,
                cond='IMC' if any(c[2]['cond'] == 'IMC' for c in ok) else 'VMC',
                vis=min(c[2]['vis_m'] for c in ok),
                techo=min(c[2]['techo_ft'] for c in ok),
                speci='SI' if any(c[0].minute != 0 for c in cl) else 'NO',
                metar=base[1],
                todos=' || '.join(r for _, r, _ in cl),
                motivo=peor[2]['motivo'],
                raros=sorted({t for _, r, _ in cl for t in raros(r)})))
        cur += timedelta(hours=1)
    return filas

# ---------------------------------------------------------------- main
ENT, SAL, CSVF = sys.argv[1], sys.argv[2], sys.argv[3]
shutil.copy(ENT, SAL)
wb = load_workbook(SAL)
wbv = load_workbook(ENT, data_only=True)   # copia con valores, para leer los totales calculados
reps = cargar(CSVF)
d1, d2 = reps[0][0].date(), reps[-1][0].date()
horas = por_hora(reps, d1, d2)
mapa = {(h['dt'].date(), h['dt'].hour): h for h in horas}

ILIM = lambda v, lim: ('ILIM' if v == lim else v) if v is not None else 's/d'
neg  = Font(bold=True, color='FFFFFF')
azul = PatternFill('solid', fgColor='1F3864')
nar  = PatternFill('solid', fgColor='F8CBAD')
gris = PatternFill('solid', fgColor='D9D9D9')
fino = Border(*[Side('thin', color='BFBFBF')] * 4)

# ---------- hoja METEO_SADM ----------
if 'METEO_SADM' in wb.sheetnames: del wb['METEO_SADM']
mt = wb.create_sheet('METEO_SADM')
mt.append(['CLAVE', 'FECHA', 'HORA_Z', 'FRANJA', 'CONDICION', 'VIS_M', 'TECHO_FT',
           'SPECI', 'METAR', 'MOTIVO', 'TODOS_LOS_REPORTES', 'TOKENS_RAROS'])
for h in horas:
    d, hh = h['dt'].date(), h['dt'].hour
    mt.append([d.toordinal() - date(1899, 12, 30).toordinal(), d, hh,
               f"{hh:02d} a {(hh+1)%24:02d}", h['cond'], ILIM(h['vis'], VIS_ILIM),
               ILIM(h['techo'], TECHO_ILIM), h['speci'], h['metar'] or 's/d', h['motivo'] or 's/d',
               h['todos'] or 's/d', ' '.join(h['raros']) or '-'])
# clave numerica = serial_fecha*100 + hora
for r in range(2, mt.max_row + 1):
    mt.cell(r, 1).value = mt.cell(r, 1).value * 100 + mt.cell(r, 3).value
for c in mt[1]: c.font, c.fill = neg, azul
mt.freeze_panes = 'A2'; mt.auto_filter.ref = mt.dimensions
for i, w in enumerate([12, 12, 8, 10, 11, 9, 10, 7, 58, 22, 80, 14], 1):
    mt.column_dimensions[get_column_letter(i)].width = w
for r in range(2, mt.max_row + 1):
    v = mt.cell(r, 5).value
    if v == 'IMC':
        for c in range(1, 13): mt.cell(r, c).fill = nar
    elif v == 'SIN DATO':
        for c in range(1, 13): mt.cell(r, c).fill = gris
mt.sheet_state = 'visible'

# ---------- columnas hora por hora en cada hoja diaria ----------
RANGO_METEO = f'METEO_SADM!$A$1:$L${mt.max_row}'   # rango acotado: columna entera es lentisimo en Google Sheets
control, tocadas = [], 0
for n in wb.sheetnames:
    if not n.isdigit(): continue
    ws = wb[n]
    f = ws['I6'].value
    fecha = f.date() if isinstance(f, datetime) else f
    ws[f'R8'] = 'CONDICION METEOROLOGICA DEL AERODROMO (SADM)'
    ws['R8'].font = Font(bold=True, size=10)
    for j, (col, cab) in enumerate(zip(COLS, CAB)):
        c = ws[f'{col}9']; c.value = cab; c.font, c.fill = neg, azul
        c.alignment = Alignment(horizontal='center', wrap_text=True)
        ws.column_dimensions[col].width = [13, 9, 11, 8, 62][j]
    for i, fila in enumerate(range(FILA0, FILAN + 1)):
        hora = i
        for j, (col, campo) in enumerate(zip(COLS, [5, 6, 7, 8, 9])):
            ws[f'{col}{fila}'] = (f'=IFERROR(VLOOKUP(INT($I$6)*100+{hora},'
                                  f'{RANGO_METEO},{campo},FALSE),"s/d")')
            ws[f'{col}{fila}'].border = fino
            ws[f'{col}{fila}'].alignment = Alignment(
                horizontal='left' if col == 'V' else 'center')
        h = mapa.get((fecha, hora)) if isinstance(fecha, date) else None
        if h and h['cond'] == 'IMC':
            for col in COLS: ws[f'{col}{fila}'].fill = nar
        elif h is None or h['cond'] == 'SIN DATO':
            for col in COLS: ws[f'{col}{fila}'].fill = gris
    tocadas += 1
    # comparacion con la marca manual (texto IMC en la grilla C..M)
    for i, fila in enumerate(range(FILA0, FILAN + 1)):
        manual = any(isinstance(ws[f'{c}{fila}'].value, str) and 'IMC' in ws[f'{c}{fila}'].value.upper()
                     for c in 'CDEFGHIJKLM')
        h = mapa.get((fecha, i)) if isinstance(fecha, date) else None
        auto = h['cond'] if h else 'FUERA DE RANGO'
        mv = wbv[n][f'H{fila}'].value
        mv = mv if isinstance(mv, (int, float)) else 0
        if (manual and auto != 'IMC') or (not manual and auto == 'IMC'):
            control.append([n, str(fecha), f"{i:02d} a {(i+1)%24:02d}", mv,
                            'IMC' if manual else '-', auto,
                            (h or {}).get('motivo', ''), (h or {}).get('metar', '')])

# ---------- hoja CONTROL ----------
if 'CONTROL' in wb.sheetnames: del wb['CONTROL']
ct = wb.create_sheet('CONTROL')
ct.append(['HOJA', 'FECHA', 'FRANJA', 'MOVIMIENTOS', 'MARCA_MANUAL', 'CLASIF_AUTOMATICA', 'MOTIVO', 'METAR'])
control.sort(key=lambda r: (-r[3], r[1], r[2]))
for f in control: ct.append(f)
for r in range(2, ct.max_row + 1):
    if (ct.cell(r, 4).value or 0) > 0:
        for c in range(1, 9): ct.cell(r, c).fill = nar
for c in ct[1]: c.font, c.fill = neg, azul
ct.freeze_panes = 'A2'; ct.auto_filter.ref = ct.dimensions
for i, w in enumerate([7, 12, 10, 13, 15, 19, 22, 62], 1):
    ct.column_dimensions[get_column_letter(i)].width = w

wb.save(SAL)
falt = sum(1 for c in control if c[4] == '-' and c[5] == 'IMC')
conmov = sum(1 for c in control if c[4] == '-' and c[5] == 'IMC' and c[3] > 0)
sobra = sum(1 for c in control if c[4] == 'IMC' and c[5] != 'IMC')
print(f"{SAL}: {tocadas} hojas diarias, {len(horas)} horas cargadas")
print(f"CONTROL: {len(control)} discrepancias -> {falt} horas IMC sin marcar ({conmov} con movimientos), {sobra} marcadas que no dan IMC")
print("METAR con tokens raros:", sorted({t for h in horas for t in h['raros']}))
